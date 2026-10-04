#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
rebuild_SI_rev5_structure.py — 以 rev4 的结构承载 rev5 的内容。

背景：`SI …_rev5.docx` 相对 rev4 有两处实质内容修改（Note S4 的
`median minimum absolute Z` 1.24 → 0.672；Table S24 表注新增 numpy.sign
零值约定句），但它的包结构被编辑器通道重写过：部件数 19 → 25，
`w:tblPrEx` 2,160 → 0，`w:tblCellMar` 2,220 → 62，`w:tblBorders` 2,222 → 62。
本项目自己的记录（`稿件归档标识更新_v4.0.2_rev6_20261004.md` 第四节）明确
弃用该通道，理由正是这组损坏。

本脚本把 rev5 的**两处文字**施加到 rev4 的**结构**上：只重写
`word/document.xml`，其余部件按 ZipInfo（同名/同压缩方式/同时间戳/同属性）
原样搬运。

用法：
    python rebuild_SI_rev5_structure.py <rev4.docx> <rev5.docx> <out.docx>
"""
import os
import re
import sys
import zipfile

# 变更 1：唯一锚点（rev4 中 "1.24" 出现两次，用整句消歧）
OLD1 = 'median minimum absolute Z = 1.24 versus'
NEW1 = 'median minimum absolute Z = 0.672 versus'

# 变更 2：在句末 run 之后插入 rev5 的新增 run
ANCHOR2 = 'no gene-level clustering applies.</w:t></w:r>'
INSERT2 = ('<w:r><w:t xml:space="preserve"> Sign agreement was computed with numpy.sign, '
           'so a statistic of exactly 0.000000 counts as disagreeing with either non-zero '
           'sign. Rounding multiZ = (wbZ + ntZ)/\u221a2 to six significant figures leaves 77 '
           'exact zeros among the 8,315 complete-case genes (wbZ 8, ntZ 7, eqZ none). Scoring '
           'zeros as positive (Z &gt; 0) instead would raise the dual-arm count from 5,506 to '
           '5,544 (+38 pairs, +0.46 percentage points); the panel-only and tissue-only arms '
           'would move by +1 and +5.</w:t></w:r>')


def doc_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read('word/document.xml').decode('utf-8')


def main(rev4, rev5, out):
    x = doc_xml(rev4)

    assert x.count(OLD1) == 1, 'anchored string 1 is not unique: %d' % x.count(OLD1)
    x = x.replace(OLD1, NEW1)

    assert x.count(ANCHOR2) == 1, 'anchored string 2 is not unique: %d' % x.count(ANCHOR2)
    x = x.replace(ANCHOR2, ANCHOR2 + INSERT2)

    with zipfile.ZipFile(rev4) as src, \
         zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'word/document.xml':
                data = x.encode('utf-8')
            info = zipfile.ZipInfo(item.filename, date_time=item.date_time)
            info.compress_type = item.compress_type
            info.external_attr = item.external_attr
            info.internal_attr = item.internal_attr
            info.create_system = item.create_system
            dst.writestr(info, data)

    # ---- report ----
    def parts(p):
        with zipfile.ZipFile(p) as z:
            return {i.filename: z.read(i.filename) for i in z.infolist()}

    a, c = parts(rev4), parts(out)
    with zipfile.ZipFile(rev5) as z:
        b = {i.filename: z.read(i.filename) for i in z.infolist()}

    print('parts        rev4 %d  ->  out %d      (rev5 was %d)' % (len(a), len(c), len(b)))
    changed = [k for k in a if c.get(k) != a[k]]
    print('rewritten parts: %s' % changed)
    y = c['word/document.xml'].decode('utf-8')
    print('tblPrEx        out %d   (rev4 %d, rev5 %d)' % (
        y.count('tblPrEx'), a['word/document.xml'].decode('utf-8').count('tblPrEx'),
        b['word/document.xml'].decode('utf-8').count('tblPrEx')))
    print('tblCellMar     out %d   (rev4 %d, rev5 %d)' % (
        y.count('tblCellMar'), a['word/document.xml'].decode('utf-8').count('tblCellMar'),
        b['word/document.xml'].decode('utf-8').count('tblCellMar')))
    print('size           %d B     (rev4 %d, rev5 %d)' % (
        os.path.getsize(out), os.path.getsize(rev4), os.path.getsize(rev5)))

    # 与 rev5 的文本比对
    import docx
    def text(p):
        d = docx.Document(p)
        o = [q.text for q in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                o.append(' | '.join(cc.text for cc in r.cells))
        return o
    if docx is not None:
        ta, tb, tc = text(rev4), text(rev5), text(out)
        print('paragraph+cells  rev4 %d  rev5 %d  out %d' % (len(ta), len(tb), len(tc)))
        print('out == rev5 (text)?   %s' % (tc == tb))
        if tc != tb:
            for i, (p, q) in enumerate(zip(tb, tc)):
                if p != q:
                    print('   first text diff at %d:\n     rev5 %r\n     out  %r'
                          % (i, p[:120], q[:120]))
                    break


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3])
