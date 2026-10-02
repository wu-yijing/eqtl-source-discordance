# -*- coding: utf-8 -*-
"""TOST 加固: 2.5 节插 TOST+功效句; Additional files 描述 S1-S8; AF1 新增 Table S8"""
from docx import Document
import shutil, os

MS = r'E:/workbuddy/BMC Genomics投稿资料/定稿资料/manuscript.docx'
AF = r'E:/workbuddy/BMC Genomics投稿资料/定稿资料/Additional file 1.docx'
TMP = r'E:/workbuddy/2026-09-09-16-50-19/audit/_tmp_fix.docx'


def repl_para(p, old, new):
    """run 级安全替换"""
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1); return True
    full = ''.join(r.text for r in p.runs)
    if old in full:
        p.runs[0].text = full.replace(old, new, 1)
        for r in p.runs[1:]:
            r.text = ''
        return True
    return False


# ========== manuscript.docx ==========
d = Document(MS)
ps = d.paragraphs

# 1) 2.5 节插入 TOST + 功效句
anchor = ('Because this comparison rests on 84 versus 81 pairs, it establishes non-difference '
          'rather than formal equivalence, and we phrase the inference accordingly.')
add = (' A pre-specified two-one-sided-test (TOST) analysis with Newcombe hybrid-score 90% CIs '
       'quantifies this resolution limit (Table S8): the observed differences are +2.9 pp (GTEx arm; '
       'housekeeping minus candidates) and \u22120.4 pp (eQTLGen arm); at a margin of \u00b115 pp, formal '
       'equivalence is satisfied in the eQTLGen arm (TOST p = 0.034) and borderline in the GTEx arm '
       '(p = 0.060), whereas a \u00b110 pp margin is not attained in either arm. The power constraint is '
       'structural: detecting a 10-pp between-group difference at 80% power would require \u2248380 '
       'gene\u2013phenotype pairs per group, beyond any curated testbed of feasible size; the present design '
       '(n = 72\u201384 pairs per group) can exclude between-group differences larger than \u224821 pp, and the '
       'one-sided 95% bounds exclude a candidate-over-control enrichment excess larger than \u224810 pp '
       '(GTEx arm) and \u224813 pp (eQTLGen arm).')
ok1 = repl_para(next(p for p in ps if anchor in p.text), anchor, anchor + add)
print('[manuscript] 2.5 TOST 句插入:', ok1)

# 2) Additional files 描述 S1-S7 -> S1-S8
ok2 = False
for p in ps:
    if 'Tables S1\u2013S7' in p.text:
        ok2 = repl_para(p, 'Tables S1\u2013S7', 'Tables S1\u2013S8 (including Table S5b)')
        break
print('[manuscript] Additional files 描述更新:', ok2)

d.save(TMP)
shutil.copyfile(TMP, MS); os.remove(TMP)
print('[manuscript] 已保存')

# ========== Additional file 1.docx: 追加 Table S8 ==========
d2 = Document(AF)
cap = ('Table S8. Two-one-sided tests (TOST) of equivalence between housekeeping-control and '
       'candidate-gene FDR enrichment rates, with Newcombe hybrid-score 90% confidence intervals.')
hdr = ['Arm (comparison)', 'Observed difference (pp)', 'Newcombe 90% CI (pp)',
       'TOST p (\u00b110 pp)', 'TOST p (\u00b115 pp)', 'TOST p (\u00b120 pp)']
rows = [
    ['GTEx arm: housekeeping 54.8% (46/84) vs candidates 51.9% (42/81)',
     '+2.9', '\u22129.7 to +15.4', '0.181', '0.060', '0.014'],
    ['eQTLGen arm: housekeeping 52.4% (44/84) vs candidates 52.8% (38/72)',
     '\u22120.4', '\u221213.3 to +12.6', '0.116', '0.034', '0.007'],
]
note = ('The observed difference is housekeeping minus candidates, in percentage points. Equivalence at '
        'margin m requires both one-sided tests significant at p < 0.05, equivalent to the 90% CI lying '
        'within (\u2212m, +m). A \u00b110 pp margin is not attained in either arm; at \u00b115 pp, equivalence is '
        'satisfied in the eQTLGen arm and borderline in the GTEx arm. Detecting a 10-pp between-group '
        'difference at 80% power would require \u2248380 gene\u2013phenotype pairs per group; the present design '
        '(n = 72\u201384 pairs per group) can exclude differences larger than \u224821 pp. The one-sided 95% '
        'bounds exclude a candidate-over-control enrichment excess larger than \u224810 pp (GTEx arm) and '
        '\u224813 pp (eQTLGen arm). See Section 2.5.')

pc = d2.add_paragraph(cap)
tbl = d2.add_table(rows=1, cols=6)
for i, h in enumerate(hdr):
    tbl.rows[0].cells[i].text = h
for row in rows:
    cs = tbl.add_row().cells
    for i, v in enumerate(row):
        cs[i].text = v
pn = d2.add_paragraph(note)

d2.save(TMP)
shutil.copyfile(TMP, AF); os.remove(TMP)
print('[AF1] Table S8 已追加到文末 (表题 + 2 数据行 + note)')
