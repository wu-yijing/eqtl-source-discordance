# -*- coding: utf-8 -*-
"""第三轮追加修订 C：HRT 受限第二重对照写入正文 + Table S11 + 数据可用性声明。"""
import os, sys, io, shutil, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt

ROOT = r"E:\workbuddy\BMC Genomics投稿资料\定稿资料"
BK   = r"E:\workbuddy\BMC Genomics投稿资料\定稿备份\备份"
TS   = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
MS   = os.path.join(ROOT, "manuscript.docx"); AF = os.path.join(ROOT, "Additional file 1.docx")
log = []
def backup(p):
    d = os.path.join(BK, os.path.basename(p).replace('.docx', '_S11修订前_%s.docx' % TS))
    shutil.copy2(p, d); log.append('backup -> ' + d)
def setp(par, txt):
    if not par.runs:
        par.add_run(txt); return
    par.runs[0].text = txt
    for r in par.runs[1:]: r.text = ''

S11_SENT = ("An HRT-restricted variant of the same control \u2014 30 genes drawn (seed 20260911) from the 818 HRT Atlas housekeeping-universe genes "
 "carrying a Whole_Blood model, again with no second-tissue requirement \u2014 gave the same result (FDR enrichment 65.3%, 49/75; 30-gene null median "
 "61.9%, 95% range 46.7\u201375.6%; Additional file 1: Table S11), so the conclusion does not depend on relaxing the housekeeping restriction, and the "
 "original HRT Atlas source file is archived with the analysis code for exact reproduction.")

# ================================================= manuscript
backup(MS)
d = Document(MS); P = d.paragraphs; n = 0
# 1) §2.2 : append the S11 sentence to the D3 paragraph
for par in P:
    if par.text.strip().startswith('To test whether this uniform enrichment reflects the control'):
        if S11_SENT not in par.text:
            setp(par, par.text + ' ' + S11_SENT)
            log.append('OK [res22-d3b appended]'); n += 1
        break
# 2) §3.6 : add Table S11 pointer
for par in P:
    if 'random genome-wide genes carrying cis-eQTL models enrich at' in par.text:
        old = '(Section 2.2; Additional file 1: Table S10).'
        new = '(Section 2.2; Additional file 1: Tables S10 and S11).'
        if old in par.text:
            setp(par, par.text.replace(old, new)); log.append('OK [lim-s11 pointer]'); n += 1
        break
# 3) additional files description
for par in P:
    if 'the genome-wide architecture-unselected random control, and the draft TWAS reporting checklist' in par.text:
        setp(par, par.text.replace('the genome-wide architecture-unselected random control, and the draft TWAS reporting checklist',
                                   'the genome-wide and HRT-restricted architecture-unselected random controls, and the draft TWAS reporting checklist'))
        log.append('OK [addfiles-s11]'); n += 1
# 4) Data availability : archive the HRT source file
for par in P:
    if par.text.strip().startswith('Processed data tables, analysis scripts and archived code: Zenodo'):
        setp(par, "Processed data tables, analysis scripts and archived code: Zenodo (concept DOI 10.5281/zenodo.21238202, which always resolves to the latest released version). The HRT Atlas v1.0 human\u2013mouse common housekeeping gene universe used for the disease-agnostic control (Methods 1.8) is archived byte-faithfully as data/hk_reselect_20260830/data/Human_Mouse_Common.csv; running the R1\u2013R5 rules of that release on it reproduces the recorded filter chain exactly (1129\u21921112\u21921013\u21921007\u2192767).")
        log.append('OK [dataavailability-hrt]'); n += 1
d.save(MS); log.append('SAVED manuscript (%d edits)' % n)

# ================================================= Additional file 1 : Table S11
backup(AF)
a = Document(AF)
cap = a.add_paragraph("Table S11. HRT-restricted, architecture-unselected random control (second control).")
rows = [
 ["Item", "Value"],
 ["Pool (HRT Atlas human\u2013mouse common HK universe \u00d7 GTEx v8 Whole_Blood MASHR model, after family- and disease-gene exclusions; no second-tissue requirement)", "818 genes (767 also carry a Nerve_Tibial model; 51 Whole_Blood-only)"],
 ["Random control (n = 30, seed = 20260911): FDR enrichment", "65.3% (49/75); p < 0.05 70.7% (53/75); p < 0.01 58.7% (44/75); p < 0.00385 53.3% (40/75)"],
 ["Null distribution (B = 10,000 draws of 30 genes from all 818): median BH enrichment (95% range)", "61.9% (46.7\u201375.6%); p < 0.05 65.4% (52.4\u201377.4%); p < 0.00385 50.0% (36.1\u201364.0%)"],
 ["Housekeeping v2 control (54.8%) percentile of BH null", "18.0th (below the random expectation)"],
 ["Candidate / non-candidate / T2DM control (58.0% / 56.8% / 59.6%) percentile of BH null", "29.2th / 25.5th / 38.1th"],
 ["Genome-wide random control (Table S10, 69.1%) percentile of this BH null", "84.0th"],
 ["Within the 818-gene pool \u2014 both-tissue vs Whole_Blood-only: BH enrichment", "64.8% (1201/1854) vs 44.4% (161/363)"],
 ["Random control gene list", "AHSA1, CALM1, CBX1, CLNS1A, COPB2, DNAJB6, GPN3, HGS, LMAN1, MLEC, NAP1L4, NUFIP2, PDS5A, POLD2, PRPF38A, PSMA6, PSMG3, PTPN11, RBM17, RNF139, SARS, SGTA, SMIM12, SOD1, STX8, TARDBP, TMEM165, UBE2Z, VDAC1, WAC"],
 ["Source file archived with the code", "data/hk_reselect_20260830/data/Human_Mouse_Common.csv (HRT Atlas v1.0; rules R1\u2013R5 reproduce 1129\u21921112\u21921013\u21921007\u2192767)"],
]
tb = a.add_table(rows=len(rows), cols=2)
try: tb.style = 'Normal Table'
except Exception as e: log.append('style warn %s' % e)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        cell = tb.cell(i, j); setp(cell.paragraphs[0], v)
        for run in cell.paragraphs[0].runs: run.font.size = Pt(8)
note = a.add_paragraph(
 "This control keeps the housekeeping restriction of Methods 1.8 but relaxes the two-tissue model requirement to Whole_Blood only, so it isolates the "
 "architecture screen within the same gene universe as the original control. Its results reproduce the genome-wide control of Table S10: random "
 "30-gene sets enrich at 61.9% (median) and the curated sets fall at or below that expectation. The HRT Atlas source file is archived byte-faithfully "
 "so that both the original control (Methods 1.8) and this second control can be regenerated exactly; running rules R1\u2013R5 of "
 "01_select_hk_genes.py on it reproduces the recorded filter chain 1129\u21921112\u21921013\u21921007\u2192767, and relaxing R5 yields this 818-gene pool.")
for run in note.runs: run.font.size = Pt(8)
a.save(AF); log.append('SAVED additional file 1 (Table S11 added)')

open(r"E:\workbuddy\2026-09-11-19-30-45\edit_log5.txt", 'w', encoding='utf-8').write("\n".join(log))
print('done')
